"""NAS 文档解析服务 —— 给 techdocs-ai 的 PDF / Word 做高精度文本提取。

手机端（pdf-lite.ts / docx-lite.ts）是自写的轻量解析器：纯 JS、离线可用、够用，
但复杂排版 PDF 会掉字、加密流处理不了。这个服务用 PyMuPDF / python-docx 做同一件事
但更准。App 侧只在「我的 → 文档解析服务」填了地址时才会调用，连不上自动退回手机本地。

接口（与 src/lib/nas-parse.ts / embedding.ts 一一对应）：
    POST /parse?name=xxx.pdf   body = 文件原始字节
        -> {"text": "...", "pages": 3, "quality": 0.97, "engine": "pymupdf"}
    POST /embed                body = {"texts": ["...", "..."]}
        -> {"vectors": [[...]], "model": "BAAI/bge-small-zh-v1.5", "dim": 512}
    GET  /health               -> {"ok": true, "engine": "...", "version": "...", "embed": {...}}

环境变量：
    AUTH_TOKEN   非空则要求 Authorization: Bearer <token>（默认空 = 不鉴权，仅限内网）
    MAX_BYTES    单个文件上限，默认 64MB
    EMBED_MODEL    嵌入模型，默认 BAAI/bge-small-zh-v1.5（中文优化 / 512 维 / CPU 够快）
    EMBED_CACHE_DIR 模型缓存目录（容器里挂 volume，避免重建容器时重下）
    EMBED_WARMUP   启动时后台预热嵌入模型，默认 1。预热失败【不影响 /parse】，
                   解析照样可用，只会在 /health 里如实报告
"""

import os
import threading
import time
import unicodedata
from contextlib import asynccontextmanager
from typing import Optional

import pymupdf  # 别再 import fitz，那条老路径已标记弃用
from fastapi import Body, FastAPI, Header, HTTPException, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse

VERSION = "1.1"
AUTH_TOKEN = os.environ.get("AUTH_TOKEN", "").strip()
MAX_BYTES = int(os.environ.get("MAX_BYTES", str(64 * 1024 * 1024)))
# 语义检索用的嵌入模型。bge-small-zh 是中文优化的小模型：512 维、CPU 上几十毫秒一条，
# 不需要 GPU 也不需要 torch（fastembed 走 ONNX），适合长期挂在 NAS 上。
EMBED_MODEL = os.environ.get("EMBED_MODEL", "BAAI/bge-small-zh-v1.5").strip()
# 模型下载缓存目录。容器里挂成 volume，否则每次重建容器都要重下（几十上百 MB）
EMBED_CACHE_DIR = os.environ.get("EMBED_CACHE_DIR", "").strip()
EMBED_WARMUP = os.environ.get("EMBED_WARMUP", "1").strip() not in ("0", "false", "")
MAX_EMBED_TEXTS = int(os.environ.get("MAX_EMBED_TEXTS", "512"))  # 单次请求最多几条
# 单条文本截断长度：bge 的上下文 512 token，中文约 1 字 1 token，超长截断即可
# （检索单元是切块后的短文本，正常不会触发）
EMBED_MAX_CHARS = int(os.environ.get("EMBED_MAX_CHARS", "1000"))

# 嵌入模型是懒加载且只加载一次的：冷启动要下模型（几十 MB），不能让 /parse 跟着一起变慢
_embedder = None
_embed_lock = threading.Lock()
_embed_error: Optional[str] = None

@asynccontextmanager
async def lifespan(_app: FastAPI):
    """后台预热嵌入模型 —— 容器一起来就加载好，免得第一次语义检索卡十几秒。

    预热失败只记录、不拦启动：解析和语义检索是两件独立的事，
    嵌入模型没就绪不该拖累 /parse。
    """
    if EMBED_WARMUP:
        threading.Thread(target=_warmup_embedder, daemon=True).start()
    yield


app = FastAPI(title="techdocs-ai parse service", version=VERSION, lifespan=lifespan)

# 手机端走原生 fetch，不受 CORS 限制；但浏览器预览（localhost:8081）需要放行。
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


def check_auth(authorization: Optional[str]) -> None:
    if not AUTH_TOKEN:
        return
    token = (authorization or "").removeprefix("Bearer ").strip()
    if token != AUTH_TOKEN:
        raise HTTPException(status_code=401, detail="令牌不正确")


def normalize_cjk_compat(text: str) -> str:
    """还原「兼容码位」汉字。

    PDF 做字体子集化时（实测 Chrome/Skia 导出、PingFang SC），常把部分汉字写成
    康熙部首区兼容码位（「用」→ U+2F64、「工」→ U+2F2F）。不还原的话用户搜
    「适用于」永远命中不了文本里的「适⽤于」。

    只对兼容区逐字 NFKC，不动全角标点 —— 与手机端 cjk-compat-table.ts 的取值范围
    严格一致，否则两条解析路径产出的文本会对不上。
    """
    out = []
    for ch in text:
        o = ord(ch)
        if 0x2E80 <= o <= 0x2FFF or 0xF900 <= o <= 0xFAFF:
            n = unicodedata.normalize("NFKC", ch)
            out.append(n if len(n) == 1 else ch)
        else:
            out.append(ch)
    return "".join(out)


def page_text_rows(page) -> str:
    """按 y 分行、按 x 拼列，输出「一行就是表格一行」的文本。

    为什么不用 page.get_text("text")：它会把同一行的每个单元格各输出一行，
    表格变成一列碎句（"性能项目 / 单位 / 数值 / 测试方法"），逐项比对时行对应关系全丢。
    这里拿 span 的 bbox 自己聚类，与手机端 pdf-lite 的输出形态保持一致。
    """
    spans = []
    for block in page.get_text("dict").get("blocks", []):
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                t = span.get("text", "")
                if t.strip():
                    b = span["bbox"]
                    spans.append((b[0], b[1], b[2], b[3], span.get("size", 10.0), t))
    if not spans:
        return ""

    spans.sort(key=lambda s: s[1])  # 先按顶边排序，再顺序聚类成行
    rows, cur = [], []
    anchor_y, row_size = None, 10.0
    for sp in spans:
        y = (sp[1] + sp[3]) / 2
        if cur and abs(y - anchor_y) <= max(sp[4], row_size) * 0.6:
            cur.append(sp)
            row_size = max(row_size, sp[4])
        else:
            if cur:
                rows.append(cur)
            cur, anchor_y, row_size = [sp], y, sp[4]
    if cur:
        rows.append(cur)

    lines = []
    for row in rows:
        row.sort(key=lambda s: s[0])
        size = max((s[3] - s[1]) for s in row)
        buf, end_x = "", None
        for x0, _y0, x1, _y1, _sz, t in row:
            # 列间距明显大于字间距时才补空格（实测表格列间隙 ≈0.65em，字间距 ≤0.28em）
            if end_x is not None and x0 - end_x > size * 0.45 and buf and not buf.endswith(" "):
                buf += " "
            buf += t
            end_x = x1
        if buf.strip():
            lines.append(buf.strip())
    return "\n".join(lines)


def score_pdf(text: str, pages: int) -> float:
    """按「每页字符数」给质量分：一页正文通常几百到几千字，个位数基本就是扫描件。"""
    if not text:
        return 0.0
    per_page = len(text) / max(pages, 1)
    if per_page < 60:
        return 0.3
    return round(min(1.0, 0.88 + per_page / 20000), 2)


def parse_pdf(data: bytes) -> dict:
    doc = pymupdf.open(stream=data, filetype="pdf")
    try:
        pages = doc.page_count
        texts = [page_text_rows(page) for page in doc]
    finally:
        doc.close()

    text = normalize_cjk_compat("\n\n".join(t for t in texts if t).strip())
    out = {
        "text": text,
        "pages": pages,
        "quality": score_pdf(text, pages),
        "engine": "pymupdf",
    }
    if not text:
        out["note"] = "PDF 中没有可提取的文本层，大概率是扫描件，需要 OCR。"
    return out


def parse_docx(data: bytes) -> dict:
    import io

    from docx import Document as DocxDocument

    d = DocxDocument(io.BytesIO(data))
    parts = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    # 表格按行拼接，单元格用空格分隔（与手机端 docx-lite 的表现保持一致）
    for tbl in d.tables:
        for row in tbl.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                parts.append("  ".join(cells))

    text = "\n\n".join(parts).strip()
    return {
        "text": text,
        "pages": 0,  # docx 无固定分页
        "quality": 0.99 if text else 0.0,
        "engine": "python-docx",
        **({} if text else {"note": "文档中没有提取到文本。"}),
    }


def get_embedder():
    """懒加载嵌入模型（进程内单例，线程安全）。首次会下载模型文件。"""
    global _embedder, _embed_error
    if _embedder is not None:
        return _embedder
    with _embed_lock:
        if _embedder is not None:
            return _embedder
        try:
            from fastembed import TextEmbedding

            kwargs = {"model_name": EMBED_MODEL}
            if EMBED_CACHE_DIR:
                kwargs["cache_dir"] = EMBED_CACHE_DIR
            _embedder = TextEmbedding(**kwargs)
            _embed_error = None
        except Exception as e:
            _embed_error = str(e)
            raise
    return _embedder


def _warmup_embedder():
    try:
        get_embedder()
        print(f"[embed] 模型就绪：{EMBED_MODEL}", flush=True)
    except Exception as e:
        print(f"[embed] 预热失败（不影响 /parse）：{e}", flush=True)


def embed_state() -> dict:
    return {"model": EMBED_MODEL, "ready": _embedder is not None, "error": _embed_error}


@app.get("/health")
def health():
    return {
        "ok": True,
        "engine": f"pymupdf {pymupdf.__version__}",
        "version": VERSION,
        "embed": embed_state(),
    }


@app.post("/embed")
async def embed(
    payload: dict = Body(default=None),
    authorization: Optional[str] = Header(None),
):
    """文本 → 向量。手机端的「语义检索」那一路靠它。

    形状与 OpenAI /embeddings 略有不同（入参 {"texts":[...]}、出参 {"vectors":[[...]]}），
    客户端 src/lib/embedding.ts 两套形状都认：填这个地址走本套，填云端走 OpenAI 那套。
    """
    check_auth(authorization)

    texts = (payload or {}).get("texts")
    if not isinstance(texts, list) or not texts:
        raise HTTPException(status_code=400, detail="texts 必须是非空字符串数组")
    if len(texts) > MAX_EMBED_TEXTS:
        raise HTTPException(status_code=413, detail=f"单次最多 {MAX_EMBED_TEXTS} 条")
    cleaned = [str(t)[:EMBED_MAX_CHARS] for t in texts]

    started = time.time()
    try:
        model = get_embedder()
        vectors = [v.tolist() for v in model.embed(cleaned)]
    except Exception as e:
        raise HTTPException(
            status_code=503,
            detail=f"嵌入模型不可用（{EMBED_MODEL}）：{e}。首次调用需要下载模型，请确认容器可联网。",
        )

    return JSONResponse({
        "vectors": vectors,
        "model": EMBED_MODEL,
        "dim": len(vectors[0]) if vectors else 0,
        "elapsedMs": int((time.time() - started) * 1000),
    })


@app.post("/parse")
async def parse(
    name: str = Query(..., description="文件名，靠扩展名判断格式"),
    authorization: Optional[str] = Header(None),
    # 默认空字节而非必填：空 body 走下面的显式 400（FastAPI 的校验转储对 App 不友好）
    payload: bytes = Body(default=b"", media_type="application/octet-stream"),
):
    check_auth(authorization)

    if len(payload) == 0:
        raise HTTPException(status_code=400, detail="请求体为空")
    if len(payload) > MAX_BYTES:
        raise HTTPException(status_code=413, detail=f"文件超过 {MAX_BYTES // 1048576}MB 上限")

    low = name.lower()
    started = time.time()
    try:
        if low.endswith(".pdf"):
            out = parse_pdf(payload)
        elif low.endswith(".docx"):
            out = parse_docx(payload)
        elif low.endswith(".doc"):
            raise HTTPException(status_code=415, detail="老式 .doc 不支持，请另存为 .docx 或 PDF")
        elif low.endswith((".txt", ".md", ".markdown")):
            out = {
                "text": payload.decode("utf-8", errors="replace").strip(),
                "pages": 0, "quality": 1.0, "engine": "raw",
            }
        else:
            raise HTTPException(status_code=415, detail=f"不支持的格式：{name}")
    except HTTPException:
        raise
    except Exception as e:  # 解析失败要给出可读原因，App 侧会展示出来
        raise HTTPException(status_code=422, detail=f"解析失败：{e}")

    out["elapsedMs"] = int((time.time() - started) * 1000)
    return JSONResponse(out)
